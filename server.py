from __future__ import annotations

import math
import os
from io import BytesIO
from pathlib import Path
from typing import Any
import joblib
import numpy as np
import pandas as pd
from fastapi import Body, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.metrics import (
    accuracy_score,
    confusion_matrix,
    f1_score,
    precision_score,
    recall_score,
)
from sklearn.model_selection import StratifiedKFold, train_test_split
from sklearn.naive_bayes import CategoricalNB
from sklearn.preprocessing import OrdinalEncoder


BASE_DIR = Path(__file__).resolve().parent
DATASET_PATH = BASE_DIR / "datasets" / "baznas-4.csv"
ARTIFACT_PATH = BASE_DIR / "artifacts" / "model_naive_bayes_baznas.pkl"

TARGET = "KELAS TARGET"
RANDOM_STATE = 42

FEATURES = [
    "JENIS USAHA/PEKERJAAN",
    "LUAS TEMPAT TINGGAL",
    "JENIS DINDING",
    "JENIS LANTAI",
    "PENDIDIKAN",
    "KATEGORI_TOTAL_PENDAPATAN",
    "KATEGORI_PENDAPATAN_PER_KAPITA",
    "KATEGORI_TANGGUNGAN",
    "KATEGORI_RIWAYAT",
]

# "KEPEMILIKAN ASET" sengaja TIDAK dijadikan fitur model (menyebabkan
# data leakage). Form web masih boleh mengirimnya; nilainya diabaikan.
RAW_CATEGORICAL_FEATURES = [
    "JENIS USAHA/PEKERJAAN",
    "LUAS TEMPAT TINGGAL",
    "JENIS DINDING",
    "JENIS LANTAI",
    "PENDIDIKAN",
]

INPUT_PAYLOAD_FIELDS = [
    *RAW_CATEGORICAL_FEATURES,
    "TOTAL_PENDAPATAN",
    "JUMLAH TANGGUNGAN",
    "RIWAYAT",
]

RAW_FIELD_ALIASES = {
    "JENIS USAHA/PEKERJAAN": ["jenis_usaha_pekerjaan", "JENIS USAHA/PEKERJAAN"],
    "LUAS TEMPAT TINGGAL": ["luas_tempat_tinggal", "LUAS TEMPAT TINGGAL"],
    "JENIS DINDING": ["jenis_dinding", "JENIS DINDING"],
    "JENIS LANTAI": ["jenis_lantai", "JENIS LANTAI"],
    "PENDIDIKAN": ["pendidikan", "PENDIDIKAN"],
}

TANGGUNGAN_NUMBER_ALIASES = ["jumlah_tanggungan", "JUMLAH TANGGUNGAN"]
TOTAL_PENDAPATAN_ALIASES = ["total_pendapatan", "TOTAL_PENDAPATAN", "TOTAL PENDAPATAN"]
RIWAYAT_ALIASES = ["riwayat", "RIWAYAT", "RIWAYAT PENERIMAAN BANTUAN"]
RIWAYAT_MAPPING = {
    0: "Belum Pernah",
    1: "1 Kali",
    2: "2 Kali",
    3: "3 Kali atau Lebih",
}

# Pilihan algoritma yang dapat dipilih dari form/API.
DEFAULT_ALGORITHM = "naive_bayes"
ALGORITHM_ALIASES = ["algoritma", "algorithm", "metode", "model", "ALGORITMA"]
ALGORITHM_LOOKUP = {
    "naive_bayes": "naive_bayes",
    "naive bayes": "naive_bayes",
    "naivebayes": "naive_bayes",
    "bayes": "naive_bayes",
    "nb": "naive_bayes",
    "c45": "c45",
    "c4.5": "c45",
    "c4_5": "c45",
    "c 4.5": "c45",
    "decision tree": "c45",
    "pohon keputusan": "c45",
}


app = FastAPI(
    title="Klasifikasi Mustahik BAZNAS",
    description="Prediksi KELAS TARGET menggunakan fitur dengan riwayat dari main.ipynb.",
    version="3.0",
)

cors_origins = [
    origin.strip()
    for origin in os.getenv("CORS_ORIGINS", "*").split(",")
    if origin.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


def kategori_pendapatan(nilai: float) -> str:
    if nilai <= 850_000:
        return "Sangat Rendah"
    if nilai <= 1_700_000:
        return "Rendah"
    if nilai <= 3_400_000:
        return "Sedang"
    return "Tinggi"


def _parse_total_pendapatan(value: Any) -> float:
    if isinstance(value, bool):
        raise HTTPException(
            status_code=422,
            detail="TOTAL_PENDAPATAN harus berupa angka, bukan boolean.",
        )

    if isinstance(value, str):
        normalized_value = value.strip().replace(".", "").replace(",", "")
    else:
        normalized_value = value

    try:
        total_pendapatan = float(normalized_value)
    except (TypeError, ValueError) as exc:
        raise HTTPException(
            status_code=422,
            detail="TOTAL_PENDAPATAN harus berupa angka.",
        ) from exc

    if not math.isfinite(total_pendapatan):
        raise HTTPException(
            status_code=422,
            detail="TOTAL_PENDAPATAN harus berupa angka valid.",
        )

    if total_pendapatan < 0:
        raise HTTPException(
            status_code=422,
            detail="TOTAL_PENDAPATAN tidak boleh negatif.",
        )

    return total_pendapatan


def kategori_tanggungan(nilai: int) -> str:
    if nilai <= 1:
        return "Rendah"
    if nilai <= 4:
        return "Sedang"
    return "Tinggi"


def kategori_riwayat(nilai: int) -> str:
    return RIWAYAT_MAPPING[nilai]


def _parse_jumlah_tanggungan(value: Any) -> int:
    if isinstance(value, bool):
        raise HTTPException(
            status_code=422,
            detail="JUMLAH TANGGUNGAN harus berupa angka, bukan boolean.",
        )

    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise HTTPException(
            status_code=422,
            detail="JUMLAH TANGGUNGAN harus berupa angka.",
        ) from exc

    if not math.isfinite(number):
        raise HTTPException(
            status_code=422,
            detail="JUMLAH TANGGUNGAN harus berupa angka valid.",
        )

    if not number.is_integer():
        raise HTTPException(
            status_code=422,
            detail="JUMLAH TANGGUNGAN harus berupa bilangan bulat.",
        )

    tanggungan = int(number)
    if tanggungan < 0:
        raise HTTPException(
            status_code=422,
            detail="JUMLAH TANGGUNGAN tidak boleh negatif.",
        )

    return tanggungan


def _parse_riwayat(value: Any) -> int:
    if isinstance(value, bool):
        raise HTTPException(
            status_code=422,
            detail="RIWAYAT harus berupa angka, bukan boolean.",
        )

    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise HTTPException(
            status_code=422,
            detail="RIWAYAT harus berupa angka.",
        ) from exc

    if not math.isfinite(number):
        raise HTTPException(
            status_code=422,
            detail="RIWAYAT harus berupa angka valid.",
        )

    if not number.is_integer():
        raise HTTPException(
            status_code=422,
            detail="RIWAYAT harus berupa bilangan bulat.",
        )

    riwayat = int(number)
    if riwayat not in RIWAYAT_MAPPING:
        raise HTTPException(
            status_code=422,
            detail="RIWAYAT harus berada pada skala 0 sampai 3.",
        )

    return riwayat


def _prepare_dataset() -> pd.DataFrame:
    if not DATASET_PATH.exists():
        raise RuntimeError(f"Dataset tidak ditemukan: {DATASET_PATH}")

    df = pd.read_csv(DATASET_PATH)
    required_columns = {
        "SUAMI",
        "ISTRI",
        "JUMLAH TANGGUNGAN",
        "RIWAYAT",
        TARGET,
        *RAW_CATEGORICAL_FEATURES,
    }
    missing = sorted(required_columns.difference(df.columns))
    if missing:
        raise RuntimeError(f"Kolom dataset tidak lengkap: {missing}")

    df["TOTAL_PENDAPATAN"] = df["SUAMI"] + df["ISTRI"]
    df["PENDAPATAN_PER_KAPITA"] = df["TOTAL_PENDAPATAN"] / (
        df["JUMLAH TANGGUNGAN"] + 1
    )
    df["KATEGORI_TOTAL_PENDAPATAN"] = df["TOTAL_PENDAPATAN"].apply(kategori_pendapatan)
    df["KATEGORI_PENDAPATAN_PER_KAPITA"] = df["PENDAPATAN_PER_KAPITA"].apply(
        kategori_pendapatan
    )
    df["KATEGORI_TANGGUNGAN"] = df["JUMLAH TANGGUNGAN"].apply(kategori_tanggungan)
    df["RIWAYAT"] = pd.to_numeric(df["RIWAYAT"], errors="coerce")
    if df["RIWAYAT"].isna().any():
        raise RuntimeError("Kolom RIWAYAT berisi nilai kosong atau bukan angka.")
    df["RIWAYAT"] = df["RIWAYAT"].clip(0, 3).astype(int)
    df["KATEGORI_RIWAYAT"] = df["RIWAYAT"].apply(kategori_riwayat)
    return df


def _normalize_text(value: Any) -> str:
    return " ".join(str(value).strip().split()).casefold()


def _is_empty_value(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    return bool(pd.isna(value))


def _build_options(df: pd.DataFrame) -> dict[str, list[str]]:
    preferred_order = {
        "KATEGORI_TOTAL_PENDAPATAN": ["Sangat Rendah", "Rendah", "Sedang", "Tinggi"],
        "KATEGORI_PENDAPATAN_PER_KAPITA": [
            "Sangat Rendah",
            "Rendah",
            "Sedang",
            "Tinggi",
        ],
        "KATEGORI_TANGGUNGAN": ["Rendah", "Sedang", "Tinggi"],
        "KATEGORI_RIWAYAT": [
            "Belum Pernah",
            "1 Kali",
            "2 Kali",
            "3 Kali atau Lebih",
        ],
    }

    options: dict[str, list[str]] = {}
    for column in FEATURES:
        values = df[column].fillna("Tidak Diketahui").astype(str).drop_duplicates().tolist()
        if column in preferred_order:
            order = preferred_order[column]
            options[column] = [value for value in order if value in values]
        else:
            options[column] = sorted(values)
    return options


def _encode(encoder: OrdinalEncoder, frame: pd.DataFrame):
    return encoder.transform(frame.fillna("Tidak Diketahui").astype(str)).astype(int) + 1


class C45Classifier:
    """Implementasi C4.5 sederhana (Gain Ratio + pre-pruning).

    Bekerja langsung pada fitur kategori (string) tanpa encoding.
    Menyediakan predict dan predict_proba agar setara dengan estimator
    scikit-learn yang dipakai pada endpoint.
    """

    def __init__(self, max_depth=5, min_samples_split=5, min_gain_ratio=0.01):
        self.max_depth = max_depth
        self.min_samples_split = min_samples_split
        self.min_gain_ratio = min_gain_ratio

    @staticmethod
    def _entropy(y: pd.Series) -> float:
        proporsi = y.value_counts(normalize=True)
        return float(-(proporsi * np.log2(proporsi)).sum())

    def _distribution(self, y: pd.Series) -> dict[str, float]:
        counts = y.value_counts()
        total = int(counts.sum())
        return {cls: float(counts.get(cls, 0)) / total for cls in self.classes_}

    def _gain_ratio(self, data: pd.DataFrame, feature: str) -> float:
        entropy_awal = self._entropy(data["_target"])
        total = len(data)
        entropy_split = 0.0
        split_info = 0.0
        for _, subset in data.groupby(feature, dropna=False):
            proporsi = len(subset) / total
            entropy_split += proporsi * self._entropy(subset["_target"])
            if proporsi > 0:
                split_info -= proporsi * np.log2(proporsi)
        information_gain = entropy_awal - entropy_split
        if split_info == 0:
            return 0.0
        return information_gain / split_info

    def _build(self, data: pd.DataFrame, features: list[str], depth: int = 0) -> dict:
        target = data["_target"]
        label = target.mode().iloc[0]
        proba = self._distribution(target)
        leaf = {"type": "leaf", "label": label, "proba": proba}

        if (
            target.nunique() == 1
            or not features
            or depth >= self.max_depth
            or len(data) < self.min_samples_split
        ):
            return leaf

        gains = {feature: self._gain_ratio(data, feature) for feature in features}
        fitur_terbaik = max(gains, key=gains.get)
        if gains[fitur_terbaik] < self.min_gain_ratio:
            return leaf

        node = {
            "type": "node",
            "feature": fitur_terbaik,
            "default": label,
            "default_proba": proba,
            "branches": {},
        }
        fitur_sisa = [feature for feature in features if feature != fitur_terbaik]
        for nilai, subset in data.groupby(fitur_terbaik, dropna=False):
            node["branches"][nilai] = self._build(
                subset.drop(columns=[fitur_terbaik]), fitur_sisa, depth + 1
            )
        return node

    def fit(self, X: pd.DataFrame, y):
        y = pd.Series(y).reset_index(drop=True)
        self.classes_ = sorted(y.unique())
        data = X.reset_index(drop=True).copy()
        data["_target"] = y
        self.tree_ = self._build(data, list(X.columns))
        return self

    def _predict_row(self, row: dict, node: dict):
        if node["type"] == "leaf":
            return node["label"], node["proba"]
        nilai = row.get(node["feature"])
        branch = node["branches"].get(nilai)
        if branch is None:
            return node["default"], node["default_proba"]
        return self._predict_row(row, branch)

    def predict(self, X: pd.DataFrame):
        return np.array(
            [self._predict_row(row, self.tree_)[0] for row in X.to_dict("records")]
        )

    def predict_proba(self, X: pd.DataFrame):
        rows = []
        for row in X.to_dict("records"):
            _, proba = self._predict_row(row, self.tree_)
            rows.append([proba[cls] for cls in self.classes_])
        return np.array(rows)


LABELS = ["TIDAK LAYAK", "LAYAK"]


def _fit_predict(algorithm: str, X_train, y_train, X_test):
    """Melatih algoritma pada satu fold/split lalu mengembalikan prediksi."""
    if algorithm == "naive_bayes":
        enc = OrdinalEncoder(handle_unknown="use_encoded_value", unknown_value=-1)
        X_train_enc = _encode(enc.fit(X_train), X_train)
        X_test_enc = _encode(enc, X_test)
        model = CategoricalNB(alpha=1.0)
        model.fit(X_train_enc, y_train)
        return model.predict(X_test_enc)

    model = C45Classifier(max_depth=5, min_samples_split=5, min_gain_ratio=0.01)
    model.fit(X_train, y_train)
    return model.predict(X_test)


def _classification_metrics(y_true, y_pred) -> dict[str, Any]:
    cm = confusion_matrix(y_true, y_pred, labels=LABELS)
    tn, fp, fn, tp = (int(value) for value in cm.ravel())
    return {
        "accuracy": float(accuracy_score(y_true, y_pred)),
        "precision_layak": float(
            precision_score(y_true, y_pred, pos_label="LAYAK", zero_division=0)
        ),
        "recall_layak": float(
            recall_score(y_true, y_pred, pos_label="LAYAK", zero_division=0)
        ),
        "f1_layak": float(
            f1_score(y_true, y_pred, pos_label="LAYAK", zero_division=0)
        ),
        "confusion_matrix": {
            "labels": LABELS,
            "matrix": cm.tolist(),
            "true_negative": tn,
            "false_positive": fp,
            "false_negative": fn,
            "true_positive": tp,
        },
    }


def _kfold_metrics(algorithm: str, X, y, n_splits: int = 5) -> dict[str, Any]:
    skf = StratifiedKFold(n_splits=n_splits, shuffle=True, random_state=RANDOM_STATE)
    per_fold: list[dict[str, Any]] = []
    accuracies: list[float] = []
    precisions: list[float] = []
    recalls: list[float] = []
    f1s: list[float] = []
    cm_total = np.zeros((2, 2), dtype=int)

    for fold, (train_index, test_index) in enumerate(skf.split(X, y), start=1):
        y_pred = _fit_predict(
            algorithm, X.iloc[train_index], y.iloc[train_index], X.iloc[test_index]
        )
        metrics = _classification_metrics(y.iloc[test_index], y_pred)
        per_fold.append({"fold": fold, **{
            key: metrics[key]
            for key in ("accuracy", "precision_layak", "recall_layak", "f1_layak")
        }})
        accuracies.append(metrics["accuracy"])
        precisions.append(metrics["precision_layak"])
        recalls.append(metrics["recall_layak"])
        f1s.append(metrics["f1_layak"])
        cm_total += np.array(metrics["confusion_matrix"]["matrix"])

    def _mean_std(values: list[float]) -> dict[str, float]:
        return {"mean": float(np.mean(values)), "std": float(np.std(values))}

    return {
        "n_splits": n_splits,
        "per_fold": per_fold,
        "ringkasan": {
            "accuracy": _mean_std(accuracies),
            "precision_layak": _mean_std(precisions),
            "recall_layak": _mean_std(recalls),
            "f1_layak": _mean_std(f1s),
        },
        "confusion_matrix_total": {
            "labels": LABELS,
            "matrix": cm_total.tolist(),
        },
    }


def _train_models(df: pd.DataFrame) -> dict[str, dict[str, Any]]:
    X = df[FEATURES].fillna("Tidak Diketahui").astype(str)
    y = df[TARGET].astype(str)

    X_train, X_test, y_train, y_test = train_test_split(
        X,
        y,
        test_size=0.20,
        stratify=y,
        random_state=RANDOM_STATE,
    )

    # ---------- Naive Bayes ----------
    encoder = OrdinalEncoder(handle_unknown="use_encoded_value", unknown_value=-1)
    X_train_encoded = _encode(encoder.fit(X_train), X_train)
    X_test_encoded = _encode(encoder, X_test)
    nb = CategoricalNB(alpha=1.0)
    nb.fit(X_train_encoded, y_train)
    nb_holdout = _classification_metrics(y_test, nb.predict(X_test_encoded))

    # ---------- C4.5 ----------
    c45 = C45Classifier(max_depth=5, min_samples_split=5, min_gain_ratio=0.01)
    c45.fit(X_train, y_train)
    c45_holdout = _classification_metrics(y_test, c45.predict(X_test))

    return {
        "naive_bayes": {
            "model": nb,
            "encoder": encoder,
            "nama": "Categorical Naive Bayes",
            "accuracy_test": nb_holdout["accuracy"],
            "source": "trained_from_dataset",
            "metrics": {
                "holdout": nb_holdout,
                "kfold": _kfold_metrics("naive_bayes", X, y),
            },
        },
        "c45": {
            "model": c45,
            "encoder": None,
            "nama": "C4.5 (Gain Ratio)",
            "accuracy_test": c45_holdout["accuracy"],
            "source": "trained_from_dataset",
            "metrics": {
                "holdout": c45_holdout,
                "kfold": _kfold_metrics("c45", X, y),
            },
        },
    }


dataset = _prepare_dataset()
options = _build_options(dataset)
option_lookup = {
    column: {_normalize_text(value): value for value in values}
    for column, values in options.items()
}
MODELS = _train_models(dataset)


def _resolve_algorithm(payload: dict[str, Any]) -> str:
    raw = None
    for alias in ALGORITHM_ALIASES:
        if alias in payload and not _is_empty_value(payload[alias]):
            raw = payload[alias]
            break
    if raw is None:
        return DEFAULT_ALGORITHM

    key = _normalize_text(raw)
    resolved = (
        ALGORITHM_LOOKUP.get(key)
        or ALGORITHM_LOOKUP.get(key.replace(" ", "_"))
        or ALGORITHM_LOOKUP.get(key.replace(" ", ""))
    )
    if resolved is None:
        raise HTTPException(
            status_code=422,
            detail=(
                f"Algoritma '{raw}' tidak dikenal. "
                "Gunakan 'Naive Bayes' atau 'C4.5'."
            ),
        )
    return resolved


def _get_payload_value(payload: dict[str, Any], feature: str) -> Any:
    for alias in RAW_FIELD_ALIASES[feature]:
        if alias in payload:
            return payload[alias]
    raise HTTPException(
        status_code=422,
        detail=f"Field '{feature}' wajib diisi.",
    )


def _get_required_numeric_payload(payload: dict[str, Any], aliases: list[str], label: str) -> Any:
    for alias in aliases:
        if alias in payload:
            value = payload[alias]
            if _is_empty_value(value):
                raise HTTPException(
                    status_code=422,
                    detail=f"{label} tidak boleh kosong.",
                )
            return value

    raise HTTPException(status_code=422, detail=f"Field '{label}' wajib diisi.")


def _validate_payload(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict) or not payload:
        raise HTTPException(status_code=422, detail="Payload JSON tidak boleh kosong.")

    total_pendapatan = _parse_total_pendapatan(
        _get_required_numeric_payload(
            payload,
            TOTAL_PENDAPATAN_ALIASES,
            "TOTAL_PENDAPATAN",
        )
    )
    jumlah_tanggungan = _parse_jumlah_tanggungan(
        _get_required_numeric_payload(
            payload,
            TANGGUNGAN_NUMBER_ALIASES,
            "JUMLAH TANGGUNGAN",
        )
    )
    riwayat = _parse_riwayat(
        _get_required_numeric_payload(
            payload,
            RIWAYAT_ALIASES,
            "RIWAYAT",
        )
    )
    pendapatan_per_kapita = total_pendapatan / (jumlah_tanggungan + 1)

    cleaned: dict[str, str] = {}
    cleaned["KATEGORI_TOTAL_PENDAPATAN"] = kategori_pendapatan(total_pendapatan)
    cleaned["KATEGORI_PENDAPATAN_PER_KAPITA"] = kategori_pendapatan(
        pendapatan_per_kapita
    )
    cleaned["KATEGORI_TANGGUNGAN"] = kategori_tanggungan(jumlah_tanggungan)
    cleaned["KATEGORI_RIWAYAT"] = kategori_riwayat(riwayat)

    for feature in RAW_CATEGORICAL_FEATURES:
        value = _get_payload_value(payload, feature)
        if _is_empty_value(value):
            raise HTTPException(
                status_code=422,
                detail=f"Field '{feature}' tidak boleh kosong.",
            )

        normalized = option_lookup[feature].get(_normalize_text(value))
        if normalized is None:
            allowed = ", ".join(options[feature])
            raise HTTPException(
                status_code=422,
                detail=(
                    f"Nilai '{value}' untuk '{feature}' tidak dikenal. "
                    f"Pilihan yang tersedia: {allowed}"
                ),
            )

        cleaned[feature] = normalized

    return {
        "features": cleaned,
        "derived_values": {
            "TOTAL_PENDAPATAN": total_pendapatan,
            "JUMLAH TANGGUNGAN": jumlah_tanggungan,
            "RIWAYAT": riwayat,
            "PENDAPATAN_PER_KAPITA": pendapatan_per_kapita,
        },
    }


def _predict_one(
    input_features: dict[str, str],
    algorithm: str = DEFAULT_ALGORITHM,
) -> dict[str, Any]:
    bundle = MODELS[algorithm]
    model = bundle["model"]

    frame = pd.DataFrame([{feature: input_features[feature] for feature in FEATURES}])
    features = _encode(bundle["encoder"], frame) if bundle["encoder"] is not None else frame

    prediction = str(model.predict(features)[0])
    probabilities = model.predict_proba(features)[0]
    probability_by_class = {
        str(label): float(probability)
        for label, probability in zip(model.classes_, probabilities)
    }

    return {
        "prediksi": prediction,
        "kelayakan": prediction,
        "confidence": float(max(probabilities)),
        "probabilitas": probability_by_class,
        "algoritma": algorithm,
        "model": {
            "nama": bundle["nama"],
            "source": bundle["source"],
            "accuracy_test": bundle["accuracy_test"],
        },
    }


def _write_excel(sheets: dict[str, pd.DataFrame]) -> BytesIO:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
    output.seek(0)
    return output


def _excel_response(buffer: BytesIO, filename: str) -> StreamingResponse:
    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _build_template_excel() -> BytesIO:
    sample = {
        "JENIS USAHA/PEKERJAAN": options["JENIS USAHA/PEKERJAAN"][0],
        "LUAS TEMPAT TINGGAL": options["LUAS TEMPAT TINGGAL"][0],
        "JENIS DINDING": options["JENIS DINDING"][0],
        "JENIS LANTAI": options["JENIS LANTAI"][0],
        "PENDIDIKAN": options["PENDIDIKAN"][0],
        "TOTAL_PENDAPATAN": 1500000,
        "JUMLAH TANGGUNGAN": 2,
        "RIWAYAT": 0,
    }
    guide_rows = []
    for column in INPUT_PAYLOAD_FIELDS:
        if column in RAW_CATEGORICAL_FEATURES:
            guide_rows.append(
                {
                    "KOLOM": column,
                    "KETERANGAN": "Wajib diisi sesuai pilihan.",
                    "CONTOH/PILIHAN": ", ".join(options[column]),
                }
            )
        elif column == "TOTAL_PENDAPATAN":
            guide_rows.append(
                {
                    "KOLOM": column,
                    "KETERANGAN": "Wajib diisi angka tanpa format rupiah.",
                    "CONTOH/PILIHAN": "1500000",
                }
            )
        elif column == "JUMLAH TANGGUNGAN":
            guide_rows.append(
                {
                    "KOLOM": column,
                    "KETERANGAN": "Wajib diisi bilangan bulat 0 atau lebih.",
                    "CONTOH/PILIHAN": "2",
                }
            )
        else:
            guide_rows.append(
                {
                    "KOLOM": column,
                    "KETERANGAN": "Wajib diisi angka skala 0 sampai 3.",
                    "CONTOH/PILIHAN": "0=Belum Pernah, 1=1 Kali, 2=2 Kali, 3=3 Kali atau Lebih",
                }
            )

    max_options = max(len(options[column]) for column in RAW_CATEGORICAL_FEATURES)
    pilihan_rows = []
    for row_index in range(max_options):
        pilihan_rows.append(
            {
                column: (
                    options[column][row_index]
                    if row_index < len(options[column])
                    else ""
                )
                for column in RAW_CATEGORICAL_FEATURES
            }
        )

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([sample], columns=INPUT_PAYLOAD_FIELDS).to_excel(
            writer,
            sheet_name="DATA",
            index=False,
        )
        pd.DataFrame(guide_rows).to_excel(writer, sheet_name="PANDUAN", index=False)
        pd.DataFrame(pilihan_rows).to_excel(writer, sheet_name="PILIHAN", index=False)

        workbook = writer.book
        data_sheet = workbook["DATA"]
        guide_sheet = workbook["PANDUAN"]
        option_sheet = workbook["PILIHAN"]

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="EDF2EB",
        )
        for sheet in [data_sheet, guide_sheet, option_sheet]:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="111E18")
                cell.fill = header_fill

        for index, column in enumerate(INPUT_PAYLOAD_FIELDS, start=1):
            data_sheet.column_dimensions[get_column_letter(index)].width = max(
                18,
                len(column) + 4,
            )
        guide_sheet.column_dimensions["A"].width = 32
        guide_sheet.column_dimensions["B"].width = 44
        guide_sheet.column_dimensions["C"].width = 80

        for index, column in enumerate(RAW_CATEGORICAL_FEATURES, start=1):
            option_sheet.column_dimensions[get_column_letter(index)].width = max(
                18,
                len(column) + 4,
            )
            data_column = get_column_letter(INPUT_PAYLOAD_FIELDS.index(column) + 1)
            option_column = get_column_letter(index)
            last_option_row = len(options[column]) + 1
            validation = DataValidation(
                type="list",
                formula1=f"'PILIHAN'!${option_column}$2:${option_column}${last_option_row}",
                allow_blank=False,
            )
            validation.error = "Pilih nilai dari dropdown agar format sesuai."
            validation.errorTitle = "Format tidak valid"
            validation.prompt = "Pilih salah satu nilai yang tersedia."
            validation.promptTitle = column
            data_sheet.add_data_validation(validation)
            validation.add(f"{data_column}2:{data_column}1000")

        numeric_validations = {
            "TOTAL_PENDAPATAN": DataValidation(
                type="decimal",
                operator="greaterThanOrEqual",
                formula1="0",
                allow_blank=False,
            ),
            "JUMLAH TANGGUNGAN": DataValidation(
                type="whole",
                operator="greaterThanOrEqual",
                formula1="0",
                allow_blank=False,
            ),
            "RIWAYAT": DataValidation(
                type="whole",
                operator="between",
                formula1="0",
                formula2="3",
                allow_blank=False,
            ),
        }
        for column, validation in numeric_validations.items():
            data_column = get_column_letter(INPUT_PAYLOAD_FIELDS.index(column) + 1)
            validation.error = "Isi angka sesuai aturan kolom."
            validation.errorTitle = "Angka tidak valid"
            data_sheet.add_data_validation(validation)
            validation.add(f"{data_column}2:{data_column}1000")

    output.seek(0)
    return output


def _format_error_detail(exc: HTTPException) -> str:
    if isinstance(exc.detail, str):
        return exc.detail
    return str(exc.detail)


def _predict_excel_rows(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.rename(columns=lambda column: str(column).strip())
    frame = frame.dropna(how="all")
    if frame.empty:
        raise HTTPException(status_code=422, detail="File Excel tidak memiliki data.")

    output_rows: list[dict[str, Any]] = []
    for row_number, row in enumerate(frame.to_dict(orient="records"), start=2):
        output_row = dict(row)
        try:
            algorithm = _resolve_algorithm(output_row)
            validated = _validate_payload(output_row)
            prediction = _predict_one(validated["features"], algorithm)
            output_row.update(
                {
                    "BARIS_EXCEL": row_number,
                    "STATUS": "success",
                    "ALGORITMA": prediction["model"]["nama"],
                    "PREDIKSI": prediction["prediksi"],
                    "CONFIDENCE": prediction["confidence"],
                    "PENDAPATAN_PER_KAPITA": validated["derived_values"][
                        "PENDAPATAN_PER_KAPITA"
                    ],
                    "KATEGORI_TOTAL_PENDAPATAN": validated["features"][
                        "KATEGORI_TOTAL_PENDAPATAN"
                    ],
                    "KATEGORI_PENDAPATAN_PER_KAPITA": validated["features"][
                        "KATEGORI_PENDAPATAN_PER_KAPITA"
                    ],
                    "KATEGORI_TANGGUNGAN": validated["features"][
                        "KATEGORI_TANGGUNGAN"
                    ],
                    "KATEGORI_RIWAYAT": validated["features"][
                        "KATEGORI_RIWAYAT"
                    ],
                    "ERROR": "",
                }
            )
            for label, probability in prediction["probabilitas"].items():
                column = "PROBABILITAS_" + str(label).upper().replace(" ", "_")
                output_row[column] = probability
        except HTTPException as exc:
            output_row.update(
                {
                    "BARIS_EXCEL": row_number,
                    "STATUS": "error",
                    "PREDIKSI": "",
                    "CONFIDENCE": "",
                    "ERROR": _format_error_detail(exc),
                }
            )
        output_rows.append(output_row)

    return pd.DataFrame(output_rows)


@app.get("/")
def home():
    return {
        "status": "API aktif",
        "endpoint": "POST /predict",
        "endpoint_excel": {
            "template": "GET /template-excel",
            "upload": "POST /predict-excel",
        },
        "fitur": FEATURES,
        "target": TARGET,
    }


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/metadata")
def metadata():
    return {
        "fitur": FEATURES,
        "input_payload": INPUT_PAYLOAD_FIELDS,
        "target": TARGET,
        "pilihan": options,
        "feature_engineering": {
            "TOTAL_PENDAPATAN": "Dikirim sebagai angka dari frontend/API.",
            "PENDAPATAN_PER_KAPITA": (
                "Dihitung otomatis: TOTAL_PENDAPATAN / (JUMLAH TANGGUNGAN + 1)"
            ),
            "KATEGORI_TOTAL_PENDAPATAN": (
                "Dibuat otomatis dari TOTAL_PENDAPATAN: Sangat Rendah jika <= 850000, "
                "Rendah jika <= 1700000, Sedang jika <= 3400000, Tinggi jika > 3400000"
            ),
            "KATEGORI_PENDAPATAN_PER_KAPITA": (
                "Dibuat otomatis dari PENDAPATAN_PER_KAPITA dengan skala kategori pendapatan yang sama"
            ),
            "KATEGORI_TANGGUNGAN": (
                "Dibuat otomatis dari JUMLAH TANGGUNGAN: "
                "Rendah jika <= 1, Sedang jika <= 4, Tinggi jika > 4"
            ),
            "RIWAYAT": "Dikirim sebagai angka skala 0 sampai 3 dari frontend/API.",
            "KATEGORI_RIWAYAT": (
                "Dibuat otomatis dari RIWAYAT: 0 Belum Pernah, 1 1 Kali, "
                "2 2 Kali, 3 3 Kali atau Lebih"
            ),
        },
        "model": {
            "default": DEFAULT_ALGORITHM,
            "tersedia": {
                key: {
                    "nama": bundle["nama"],
                    "source": bundle["source"],
                    "accuracy_test": bundle["accuracy_test"],
                }
                for key, bundle in MODELS.items()
            },
        },
    }


@app.get("/metrics")
def metrics_all():
    return {
        "target": TARGET,
        "labels": LABELS,
        "default": DEFAULT_ALGORITHM,
        "algoritma": {
            key: {
                "nama": bundle["nama"],
                "accuracy_test": bundle["accuracy_test"],
                "holdout": bundle["metrics"]["holdout"],
                "kfold": bundle["metrics"]["kfold"],
            }
            for key, bundle in MODELS.items()
        },
    }


@app.get("/metrics/{algorithm}")
def metrics_one(algorithm: str):
    key = _resolve_algorithm({"algoritma": algorithm})
    bundle = MODELS[key]
    return {
        "algoritma": key,
        "nama": bundle["nama"],
        "accuracy_test": bundle["accuracy_test"],
        "target": TARGET,
        "labels": LABELS,
        "holdout": bundle["metrics"]["holdout"],
        "kfold": bundle["metrics"]["kfold"],
    }


@app.get("/template-excel")
@app.get("/download-template-excel")
@app.get("/format-excel")
@app.get("/template_excel")
def download_template_excel():
    return _excel_response(_build_template_excel(), "format-prediksi-baznas.xlsx")


@app.post("/predict-excel")
@app.post("/upload-excel")
@app.post("/predict_excel")
async def predict_excel(file: UploadFile = File(...)):
    filename = file.filename or ""
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=422,
            detail="File harus berformat Excel .xlsx.",
        )

    content = await file.read()
    if not content:
        raise HTTPException(status_code=422, detail="File Excel tidak boleh kosong.")

    try:
        frame = pd.read_excel(BytesIO(content), sheet_name=0)
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="File Excel gagal dibaca. Pastikan memakai format template.",
        ) from exc

    result_frame = _predict_excel_rows(frame)
    result_excel = _write_excel({"HASIL_PREDIKSI": result_frame})
    return _excel_response(result_excel, "hasil-prediksi-baznas.xlsx")


@app.post("/predict")
def predict(payload: dict[str, Any] = Body(...)):
    try:
        algorithm = _resolve_algorithm(payload)
        validated = _validate_payload(payload)
        input_features = validated["features"]
        result = _predict_one(input_features, algorithm)
        return {
            "status": "success",
            "target": TARGET,
            "fitur_digunakan": FEATURES,
            "input": input_features,
            "nilai_hasil_feature_engineering": validated["derived_values"],
            "feature_engineering": {
                "KATEGORI_TOTAL_PENDAPATAN": (
                    "Dihitung otomatis dari TOTAL_PENDAPATAN berdasarkan skala pendapatan"
                ),
                "KATEGORI_PENDAPATAN_PER_KAPITA": (
                    "Dihitung otomatis dari TOTAL_PENDAPATAN / (JUMLAH TANGGUNGAN + 1)"
                ),
                "KATEGORI_TANGGUNGAN": (
                    "Rendah jika JUMLAH TANGGUNGAN <= 1, "
                    "Sedang jika <= 4, Tinggi jika > 4"
                ),
                "KATEGORI_RIWAYAT": (
                    "0 Belum Pernah, 1 1 Kali, 2 2 Kali, 3 3 Kali atau Lebih"
                ),
            },
            **result,
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        "server:app",
        host=os.getenv("HOST", "0.0.0.0"),
        port=int(os.getenv("PORT", "8000")),
        reload=os.getenv("RELOAD", "false").lower() == "true",
    )
